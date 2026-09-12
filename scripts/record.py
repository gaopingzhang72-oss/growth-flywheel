# -*- coding: utf-8 -*-
"""
复盘记录写入器 —— 个人成长飞轮的每日/每周复盘表格记账

用法（Agent 用）：
  python record.py init                                   # 初始化表格（已存在则不动）
  python record.py daily  --json 条目.json                # 追加/覆盖某一天的一行
  python record.py weekly --json 条目.json                # 追加/覆盖某一周的一行
  python record.py stats                                  # 只重算「统计」表

条目.json 的字段（daily）：
  date      "2026-09-13"（可省略，默认今天）
  summary   今天学了什么（摘要）
  blocker   卡在哪
  tomorrow  明天第一件事
  lines     动了哪几条线，字符串或数组，例如 "C语言,打字"
  sport     今天运动了什么 + 多久，例如 "跑步 30 分钟"；没运动写 "无"（可选）
  score     自评 1~5（可选）
  raw       原始汇报摘录（可选，<=500 字）
  note      备注（可选）

条目.json 的字段（weekly）：
  date      该周任意一天（默认今天），用来自动算周次和区间
  grade     绩点/课内一句
  clab      C语言+实验室一句
  english   英语一句
  sport     运动一句（本周动了几次、有没有动起来）（可选）
  problem   本周主要问题
  goals     下周三个目标
  last      上周目标达成情况（回滚回填）
  extra     备注（可选）

设计原则：只追加、同一天/同一周重复写入时覆盖那一行，绝不删除历史。
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# 台账位置：优先环境变量 REVIEW_XLSX，否则默认放在本脚本所在目录（tools/）的上一级
_HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_FILE = os.environ.get("REVIEW_XLSX") or os.path.normpath(
    os.path.join(_HERE, os.pardir, "每日复盘记录.xlsx"))

DAILY_SHEET, WEEKLY_SHEET, STATS_SHEET = "① 每日打卡", "② 每周复盘", "③ 统计"

DAILY_HDR = ["日期", "星期", "今天学了什么", "卡在哪", "明天第一件事",
             "动了哪几条线", "运动", "自评(1-5)", "原始汇报摘录", "备注"]
WEEKLY_HDR = ["周次", "日期区间", "绩点/课内", "C语言+实验室", "英语", "运动",
              "本周主要问题", "下周三个目标", "上周目标达成情况", "打卡天数", "备注"]

# 列宽
DAILY_W = [12, 6, 42, 34, 30, 18, 22, 10, 46, 20]
WEEKLY_W = [10, 18, 22, 24, 22, 24, 34, 40, 34, 10, 18]

# 列号按名字算，以后插列不用再改一堆数字
D_DATE = DAILY_HDR.index("日期") + 1
D_LINES = DAILY_HDR.index("动了哪几条线") + 1
D_SPORT = DAILY_HDR.index("运动") + 1
D_SCORE = DAILY_HDR.index("自评(1-5)") + 1
D_RAW = DAILY_HDR.index("原始汇报摘录") + 1

LINES = ["C语言", "打字", "英语", "实验室", "课内"]
WEEKDAYS = ["一", "二", "三", "四", "五", "六", "日"]

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=11)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_sheet(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def init(path):
    if os.path.exists(path):
        print("[init] 表格已存在，未改动：", path)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = DAILY_SHEET
    _style_sheet(ws, DAILY_HDR, DAILY_W)
    ws2 = wb.create_sheet(WEEKLY_SHEET)
    _style_sheet(ws2, WEEKLY_HDR, WEEKLY_W)
    wb.create_sheet(STATS_SHEET)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    rebuild_stats(path)
    print("[init] 已创建：", path)


def _headers(ws):
    return [str(ws.cell(row=1, column=i).value or "").strip()
            for i in range(1, ws.max_column + 1)]


def _ensure_new_columns(ws, headers, widths):
    """老表格缺少新列时补上：插在和表头定义相同的位置，老数据自动右移，不丢内容。"""
    for i, h in enumerate(headers):
        cur = _headers(ws)
        if h in cur:
            continue
        if i + 1 <= ws.max_column and cur[i:i + 1]:
            ws.insert_cols(i + 1)          # 老列整体右移
        c = ws.cell(row=1, column=i + 1, value=h)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(i + 1)].width = widths[i]
        print(f"[schema] 已补新列：{h}")


def _load(path):
    if not os.path.exists(path):
        init(path)
    wb = load_workbook(path)
    if DAILY_SHEET in wb.sheetnames:
        _ensure_new_columns(wb[DAILY_SHEET], DAILY_HDR, DAILY_W)
    if WEEKLY_SHEET in wb.sheetnames:
        _ensure_new_columns(wb[WEEKLY_SHEET], WEEKLY_HDR, WEEKLY_W)
    return wb


def _upsert(ws, key_col, key_val, values, ncols):
    """同 key 覆盖，否则追加到末尾。返回 ("updated"/"appended", 行号)"""
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=key_col).value).strip() == str(key_val).strip():
            for i, v in enumerate(values, start=1):
                ws.cell(row=r, column=i, value=v)
            return "updated", r
    r = ws.max_row + 1
    if r == 2 and ws.cell(row=2, column=1).value is None:
        r = 2
    for i, v in enumerate(values, start=1):
        ws.cell(row=r, column=i, value=v)
    return "appended", r


def _prettify(ws, ncols_wide, wrap_cols, center_cols=(1, 2, D_SCORE)):
    for r in range(2, ws.max_row + 1):
        for c in range(1, ncols_wide + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in wrap_cols),
                                       horizontal="center" if c in center_cols else "left")


def _parse_date(s):
    if not s:
        return dt.date.today()
    s = str(s).strip().replace("/", "-").replace(".", "-")
    for f in ("%Y-%m-%d", "%Y-%m-%d %H:%M", "%y-%m-%d"):
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    print("[warn] 日期无法解析，用今天代替：", s)
    return dt.date.today()


def daily(path, item):
    d = _parse_date(item.get("date"))
    lines = item.get("lines") or ""
    if isinstance(lines, (list, tuple)):
        lines = "、".join(str(x) for x in lines)
    raw = (item.get("raw") or "").strip()
    if len(raw) > 500:
        raw = raw[:500] + "…"
    sport = (item.get("sport") or "").strip()
    values = [d.isoformat(), WEEKDAYS[d.weekday()],
              (item.get("summary") or "").strip(),
              (item.get("blocker") or "").strip(),
              (item.get("tomorrow") or "").strip(),
              lines, sport, item.get("score") or "", raw,
              (item.get("note") or "").strip()]
    wb = _load(path)
    ws = wb[DAILY_SHEET]
    action, r = _upsert(ws, 1, d.isoformat(), values, len(DAILY_HDR))
    _prettify(ws, len(DAILY_HDR), {3, 4, 5, 9}, (1, 2, D_SCORE))
    wb.save(path)
    rebuild_stats(path)
    print(f"[daily] {action} 第 {r} 行：{d.isoformat()} 动线={lines or '未填'} "
          f"运动={sport or '未填'}")


def weekly(path, item):
    d = _parse_date(item.get("date"))
    iso = d.isocalendar()
    week_no = item.get("week") or f"第{iso[1]}周"
    monday = d - dt.timedelta(days=d.weekday())
    sunday = monday + dt.timedelta(days=6)
    rng = item.get("range") or f"{monday.isoformat()} ~ {sunday.isoformat()}"

    wb = _load(path)
    ws = wb[DAILY_SHEET]
    days = 0
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v and _parse_date(v).isocalendar()[1] == iso[1]:
            days += 1
    ws2 = wb[WEEKLY_SHEET]
    values = [week_no, rng,
              (item.get("grade") or "").strip(), (item.get("clab") or "").strip(),
              (item.get("english") or "").strip(), (item.get("sport") or "").strip(),
              (item.get("problem") or "").strip(),
              (item.get("goals") or "").strip(), (item.get("last") or "").strip(),
              days, (item.get("extra") or "").strip()]
    action, r = _upsert(ws2, 1, week_no, values, len(WEEKLY_HDR))
    _prettify(ws2, len(WEEKLY_HDR), {2, 3, 4, 5, 6, 7, 8, 9}, (1, 10))
    wb.save(path)
    rebuild_stats(path)
    print(f"[weekly] {action} {week_no}（{rng}）本周打卡 {days} 天")


def _all_dates(ws):
    out = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=D_DATE).value
        if v:
            out.append(_parse_date(v))
    return sorted(out)


NO_SPORT = re.compile(r"^(无|没有|没|否|没运动|没练|未运动|未动|休息|0|[-—–、/\s]*)$")


def _has_sport(text):
    """这一行的「运动」算不算真动了。"""
    s = str(text or "").strip()
    if not s or NO_SPORT.match(s):
        return False
    return True


def _sport_minutes(text):
    """从「跑步 30 分钟」「游泳 1.5 小时」里抠出分钟数，抠不到就算 0。"""
    s = str(text or "")
    total = 0.0
    for x in re.findall(r"(\d+(?:\.\d+)?)\s*(?:小时|个小时|h|hr|hours?|hour)",
                        s, flags=re.I):
        total += float(x) * 60
    rest = re.sub(r"\d+(?:\.\d+)?\s*(?:小时|个小时|h|hr|hours?|hour)s?", "", s,
                  flags=re.I)
    for x in re.findall(r"(\d+(?:\.\d+)?)\s*(?:分钟|分|min|mins|minutes|m)", rest,
                        flags=re.I):
        total += float(x)
    return int(round(total))


def rebuild_stats(path):
    wb = _load(path)
    ws = wb[DAILY_SHEET]
    wsu = wb[WEEKLY_SHEET]
    if STATS_SHEET in wb.sheetnames:
        del wb[STATS_SHEET]
    st = wb.create_sheet(STATS_SHEET)
    st.column_dimensions["A"].width = 24
    st.column_dimensions["B"].width = 14
    for col in "CDE":
        st.column_dimensions[col].width = 12

    dates = _all_dates(ws)
    rows_ok = [r for r in range(2, ws.max_row + 1)
               if ws.cell(row=r, column=D_DATE).value]
    scores = [ws.cell(row=r, column=D_SCORE).value for r in rows_ok]
    scores = [s for s in scores if isinstance(s, (int, float))]
    sport_days = sum(1 for r in rows_ok if _has_sport(ws.cell(row=r, column=D_SPORT).value))
    sport_min = sum(_sport_minutes(ws.cell(row=r, column=D_SPORT).value) for r in rows_ok)

    streak = 0
    if dates:
        cur = dt.date.today()
        pool = set(dates)
        while cur in pool:
            streak += 1
            cur -= dt.timedelta(days=1)
        if streak == 0 and (cur := dt.date.today() - dt.timedelta(days=1)) in pool:
            while cur in pool:
                streak += 1
                cur -= dt.timedelta(days=1)

    st["A1"] = "复盘统计（每次写入后自动重算）"
    st["A1"].font = Font(name="微软雅黑", size=13, bold=True, color="1F4E79")
    rows = [
        ("总打卡天数", len(dates)),
        ("连续打卡天数", streak),
        ("最近一次打卡", dates[-1].isoformat() if dates else "—"),
        ("平均自评", round(sum(scores) / len(scores), 2) if scores else "—"),
        ("写成周复盘的周数", sum(1 for r in range(2, wsu.max_row + 1) if wsu.cell(row=r, column=1).value)),
        ("有运动的打卡天数", sport_days),
        ("累计运动时长(分钟)", sport_min),
    ]
    r = 3
    for k, v in rows:
        st.cell(row=r, column=1, value=k).font = Font(name="微软雅黑", size=11, bold=True)
        c = st.cell(row=r, column=2, value=v)
        c.font = BODY_FONT
        c.alignment = Alignment(horizontal="center")
        r += 1
    r += 1

    st.cell(row=r, column=1, value="各条线出现次数").font = Font(name="微软雅黑", size=12, bold=True, color="1F4E79")
    r += 1
    for ln in LINES:
        n = sum(1 for row in rows_ok
                if ln in str(ws.cell(row=row, column=D_LINES).value or ""))
        st.cell(row=r, column=1, value=ln).font = BODY_FONT
        c = st.cell(row=r, column=2, value=n)
        c.font = BODY_FONT
        c.alignment = Alignment(horizontal="center")
        r += 1
    r += 1

    # 图表数据：按周汇总
    base = r
    heads = ["周次", "打卡天数", "平均自评", "运动天数"]
    for i, h in enumerate(heads, start=1):
        c = st.cell(row=base, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center")

    by_week = {}
    for row in rows_ok:
        d = _parse_date(ws.cell(row=row, column=D_DATE).value)
        wk = d.isocalendar()[1]
        by_week.setdefault(wk, {"n": 0, "s": [], "sp": 0})
        by_week[wk]["n"] += 1
        s = ws.cell(row=row, column=D_SCORE).value
        if isinstance(s, (int, float)):
            by_week[wk]["s"].append(s)
        if _has_sport(ws.cell(row=row, column=D_SPORT).value):
            by_week[wk]["sp"] += 1

    rr = base + 1
    for wk in sorted(by_week):
        st.cell(row=rr, column=1, value=f"第{wk}周").font = BODY_FONT
        st.cell(row=rr, column=2, value=by_week[wk]["n"]).font = BODY_FONT
        avg = (round(sum(by_week[wk]["s"]) / len(by_week[wk]["s"]), 2)
               if by_week[wk]["s"] else None)
        st.cell(row=rr, column=3, value=avg if avg is not None else "").font = BODY_FONT
        st.cell(row=rr, column=4, value=by_week[wk]["sp"]).font = BODY_FONT
        rr += 1

    if rr > base + 1:
        ch = LineChart()
        ch.title = "每周打卡天数 / 平均自评 / 运动天数"
        ch.height, ch.width = 8, 18
        data = Reference(st, min_col=2, max_col=4, min_row=base, max_row=rr - 1)
        cats = Reference(st, min_col=1, min_row=base + 1, max_row=rr - 1)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        st.add_chart(ch, "F3")

    wb.save(path)
    return {"days": len(dates), "streak": streak, "sport_days": sport_days,
            "sport_minutes": sport_min}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["init", "daily", "weekly", "stats"])
    ap.add_argument("--file", default=DEFAULT_FILE)
    ap.add_argument("--json", dest="json_path")
    a = ap.parse_args()

    if a.cmd == "init":
        return init(a.file)
    if a.cmd == "stats":
        print("[stats]", rebuild_stats(a.file))
        return
    item = {}
    if a.json_path:
        with open(a.json_path, encoding="utf-8") as f:
            item = json.load(f)
    if a.cmd == "daily":
        daily(a.file, item)
    else:
        weekly(a.file, item)


if __name__ == "__main__":
    main()
